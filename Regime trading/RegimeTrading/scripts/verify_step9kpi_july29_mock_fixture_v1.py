from __future__ import annotations

import argparse
import json
import math
import tempfile
from pathlib import Path

from openpyxl import load_workbook

from RegimeTrading.scripts.step9kpi_read_only_evaluation_v1 import build


SESSION = "2026-07-29"


def _one(frame, **filters):
    selected = frame.copy()
    for column, value in filters.items():
        selected = selected.loc[selected[column].astype(str).eq(str(value))]
    if len(selected) != 1:
        raise AssertionError(f"Expected exactly one row for {filters}; found {len(selected)}")
    return selected.iloc[0]


def _close(actual, expected, tolerance=1e-6):
    if actual is None or not math.isfinite(float(actual)) or abs(float(actual) - expected) > tolerance:
        raise AssertionError(f"Expected {expected} +/- {tolerance}; found {actual}")


def verify(project_root: Path, config: Path, schema: Path) -> None:
    project_root = project_root.resolve()
    with tempfile.TemporaryDirectory(prefix="step9kpi_mock_fixture_") as temp:
        out = Path(temp) / "step9kpi"
        workbook = out / "powerbi_step9_kpi_monitor.xlsx"
        result = build(
            project_root=project_root,
            config_path=config,
            schema_path=schema,
            output_dir=out,
            workbook_path=workbook,
            session_dates=[SESSION],
        )

        if result.source_hashes_before != result.source_hashes_after:
            raise AssertionError("Source ledgers changed during KPI fixture verification.")

        engine = result.tables["tblEngineDaily"]
        benchmark = result.tables["tblBenchmarkDaily"]
        regime = result.tables["tblRegimeAccuracy"]
        ranking = result.tables["tblRankingDaily"]
        quality = result.tables["tblDataQuality"]

        _close(_one(engine, session_date=SESSION, engine_book_id="STEP9L_V3_SELECTED")["native_net_pnl_sek"], 3.092063835681593)
        _close(_one(engine, session_date=SESSION, engine_book_id="STEP9S_MANDATORY_CONTROL")["native_net_pnl_sek"], 3.950943211522335)
        _close(_one(engine, session_date=SESSION, engine_book_id="STEP9R_SELECTED")["native_net_pnl_sek"], 0.0)
        _close(_one(engine, session_date=SESSION, engine_book_id="STEP9U_SELECTED")["native_net_pnl_sek"], 7.453894915379179)

        oracle = _one(benchmark, session_date=SESSION, benchmark_id="ORACLE_TOP2_OBSERVED_FIXED")
        if oracle["selected_tickers"] != "ABB.ST|ATCO-A.ST":
            raise AssertionError(f"Unexpected unrestricted oracle tickers: {oracle['selected_tickers']}")
        _close(oracle["standardized_net_pnl_sek"], 28.160371332132893)

        capped_oracle = _one(
            benchmark,
            session_date=SESSION,
            benchmark_id="ORACLE_TOP2_OBSERVED_SECTOR_CAPPED_FIXED",
        )
        if capped_oracle["selected_tickers"] != "ABB.ST|GETI-B.ST":
            raise AssertionError(f"Unexpected sector-capped oracle tickers: {capped_oracle['selected_tickers']}")
        _close(capped_oracle["standardized_net_pnl_sek"], 27.702846147908634)

        primary_regime = _one(regime, session_date=SESSION, classifier_universe="REGIME_SOURCE")
        if primary_regime["morning_regime"] != "RANGE_LOW_VOL":
            raise AssertionError(f"Unexpected morning regime: {primary_regime['morning_regime']}")
        if primary_regime["realized_eod_regime"] != "DEFENSIVE_MIXED":
            raise AssertionError(f"Unexpected realized EOD regime: {primary_regime['realized_eod_regime']}")
        if primary_regime["regime_accuracy_state"] != "MISMATCH":
            raise AssertionError(f"Unexpected regime accuracy state: {primary_regime['regime_accuracy_state']}")

        step9r = _one(ranking, session_date=SESSION, engine_book_id="STEP9R_SELECTED")
        _close(step9r["total_selection_opportunity_loss_sek"], 4.122751780908791)
        step9u = _one(ranking, session_date=SESSION, engine_book_id="STEP9U_SELECTED")
        _close(step9u["total_selection_opportunity_loss_sek"], 4.806632032828186)

        failed_quality = quality.loc[~quality["passed"].fillna(False).astype(bool)]
        if not failed_quality.empty:
            raise AssertionError(
                "Data-quality checks failed: "
                + json.dumps(failed_quality[["check_id", "source_id", "detail"]].to_dict("records"), default=str)
            )

        book = load_workbook(workbook, read_only=False, data_only=False)
        expected_tables = set(result.tables)
        actual_tables = {
            table_name
            for sheet in book.worksheets
            for table_name in sheet.tables.keys()
        }
        if expected_tables != actual_tables:
            raise AssertionError(f"Workbook named tables mismatch. Expected={sorted(expected_tables)} actual={sorted(actual_tables)}")

        manifest = json.loads((out / "step9kpi_build_manifest.json").read_text(encoding="utf-8"))
        if manifest.get("router_active") is not False or manifest.get("orders_enabled") is not False:
            raise AssertionError("Router/order safety contract failed in manifest.")
        if manifest.get("source_files_unchanged") is not True:
            raise AssertionError("Manifest did not confirm unchanged source files.")

    print("STEP9KPI_JULY29_MOCK_FIXTURE_VERIFICATION: PASSED")
    print("SESSION: 2026-07-29")
    print("EVIDENCE STATUS: MOCK_REHEARSAL / NON_CONFIRMATORY")
    print("STEP 9L NATIVE P&L: +3.092064 SEK")
    print("STEP 9S CONTROL NATIVE P&L: +3.950943 SEK")
    print("STEP 9R SELECTED P&L: 0.000000 SEK")
    print("STEP 9U NATIVE P&L: +7.453895 SEK")
    print("UNRESTRICTED ORACLE TOP 2: ABB.ST | ATCO-A.ST")
    print("UNRESTRICTED ORACLE TOP 2 STANDARDIZED P&L: +28.160371 SEK")
    print("SECTOR-CAPPED ORACLE TOP 2: ABB.ST | GETI-B.ST")
    print("SECTOR-CAPPED ORACLE TOP 2 STANDARDIZED P&L: +27.702846 SEK")
    print("MORNING / REALIZED REGIME: RANGE_LOW_VOL / DEFENSIVE_MIXED")
    print("SOURCE LEDGERS: BYTE-FOR-BYTE UNCHANGED")
    print("ROUTER ACTIVE: FALSE")
    print("NO ORDER WAS SENT")


def _parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Verify the approved July 29 Step 9 KPI mock fixtures.")
    parser.add_argument("--project-root", type=Path, required=True)
    parser.add_argument("--config", type=Path, required=True)
    parser.add_argument("--schema", type=Path, required=True)
    return parser.parse_args()


def main() -> int:
    args = _parse_args()
    verify(args.project_root, args.config, args.schema)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
